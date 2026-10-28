"""Motores de Gajo. Cada operacion recibe rutas y devuelve las rutas que creo.

Regla de oro: nunca se modifica ni se borra el archivo original; el resultado
queda al lado, con un nombre libre.
"""
from __future__ import annotations

import io
import logging
import math
import re
import shutil
import tarfile
import zipfile
from pathlib import Path

from .tools import (
    BSDTAR, GajoError, base_stem, ffmpeg, norm_ext, pandoc, place, probe, run,
    soffice_convert, unique_dir, unique_path, work_dir,
)

log = logging.getLogger("gajo")

# ---------------------------------------------------------------- imagenes

_PIL_FMT = {"jpg": "JPEG", "png": "PNG", "webp": "WEBP", "heic": "HEIF", "tiff": "TIFF",
            "avif": "AVIF", "bmp": "BMP", "pdf": "PDF", "gif": "GIF"}
_heif_ready = False


def _heif() -> None:
    global _heif_ready
    if not _heif_ready:
        import pillow_heif
        pillow_heif.register_heif_opener()
        _heif_ready = True


def open_image(path: Path):
    from PIL import Image, ImageOps
    _heif()
    if norm_ext(path) == "svg":
        import fitz
        with fitz.open(path) as doc:
            pix = doc[0].get_pixmap(dpi=192, alpha=True)
            return Image.frombytes("RGBA", (pix.width, pix.height), pix.samples)
    im = Image.open(path)
    im.load()
    return ImageOps.exif_transpose(im)


def _has_alpha(im) -> bool:
    return im.mode in ("RGBA", "LA", "PA") or (im.mode == "P" and "transparency" in im.info)


def _flatten(im):
    from PIL import Image
    if _has_alpha(im):
        rgba = im.convert("RGBA")
        bg = Image.new("RGB", rgba.size, (255, 255, 255))
        bg.paste(rgba, mask=rgba.getchannel("A"))
        return bg
    return im if im.mode in ("RGB", "L") else im.convert("RGB")


def save_image(im, out: Path, fmt: str, quality: int | None = None, keep_meta: bool = True) -> None:
    exif = im.info.get("exif") if keep_meta else None
    icc = im.info.get("icc_profile")
    if fmt in ("jpg", "bmp", "pdf"):
        im = _flatten(im)
    elif fmt in ("heic", "avif", "webp"):
        im = im.convert("RGBA" if _has_alpha(im) else "RGB")
    elif im.mode not in ("RGB", "RGBA", "L", "LA", "P", "1"):
        im = im.convert("RGBA" if _has_alpha(im) else "RGB")
    kw: dict = {}
    if fmt == "jpg":
        kw = {"quality": quality or 92, "optimize": True, "progressive": True}
    elif fmt == "webp":
        kw = {"quality": quality or 90, "method": 4}
    elif fmt == "heic":
        kw = {"quality": quality or 90}
    elif fmt == "avif":
        kw = {"quality": quality or 80}
    elif fmt == "tiff":
        kw = {"compression": "tiff_lzw"}
    elif fmt == "png":
        kw = {"optimize": True}
    elif fmt == "pdf":
        kw = {"resolution": 150}
    if exif and fmt in ("jpg", "webp", "heic", "avif", "png", "tiff"):
        kw["exif"] = exif
    if icc and fmt not in ("bmp", "gif"):
        kw["icc_profile"] = icc
    im.save(out, _PIL_FMT[fmt], **kw)


def img_convert(src: Path, fmt: str) -> list[Path]:
    if norm_ext(src) == "gif" and fmt == "mp4":
        out = unique_path(src.parent, src.stem, "mp4")
        ffmpeg(["-i", src, "-movflags", "+faststart", "-pix_fmt", "yuv420p",
                "-vf", "scale=trunc(iw/2)*2:trunc(ih/2)*2", out])
        return [out]
    im = open_image(src)
    out = unique_path(src.parent, src.stem, fmt)
    save_image(im, out, fmt)
    return [out]


def _same_fmt(src: Path) -> str:
    """Formato de salida para herramientas que conservan el tipo de imagen."""
    ext = norm_ext(src)
    return ext if ext in ("jpg", "png", "webp", "heic", "avif") else "png"


def _img_tool(src: Path, suffix: str, fn, fmt: str | None = None, **save_kw) -> list[Path]:
    im = fn(open_image(src))
    fmt = fmt or _same_fmt(src)
    out = unique_path(src.parent, f"{src.stem}-{suffix}", fmt)
    save_image(im, out, fmt, **save_kw)
    return [out]


def img_compress(src: Path) -> list[Path]:
    from PIL import Image
    ext = norm_ext(src)
    im = open_image(src)
    im.thumbnail((2560, 2560), Image.Resampling.LANCZOS)
    if ext == "png":
        quant = im.convert("RGBA").quantize(colors=256, method=Image.Quantize.FASTOCTREE,
                                            dither=Image.Dither.FLOYDSTEINBERG)
        out = unique_path(src.parent, f"{src.stem}-comprimido", "png")
        quant.save(out, "PNG", optimize=True)
        return [out]
    fmt = ext if ext in ("jpg", "webp", "heic", "avif") else ("png" if _has_alpha(im) else "jpg")
    quality = {"jpg": 72, "webp": 70, "heic": 60, "avif": 55}.get(fmt)
    out = unique_path(src.parent, f"{src.stem}-comprimido", fmt)
    save_image(im, out, fmt, quality=quality)
    return [out]


def img_half(src: Path) -> list[Path]:
    from PIL import Image
    return _img_tool(src, "50", lambda im: im.resize(
        (max(1, im.width // 2), max(1, im.height // 2)), Image.Resampling.LANCZOS))


def img_rotate(src: Path) -> list[Path]:
    return _img_tool(src, "girada", lambda im: im.rotate(-90, expand=True))


def img_square(src: Path) -> list[Path]:
    def crop(im):
        side = min(im.size)
        left, top = (im.width - side) // 2, (im.height - side) // 2
        return im.crop((left, top, left + side, top + side))
    return _img_tool(src, "cuadrada", crop)


def img_gray(src: Path) -> list[Path]:
    from PIL import ImageOps
    return _img_tool(src, "byn", lambda im: ImageOps.grayscale(_flatten(im)).convert("RGB"))


def img_strip(src: Path) -> list[Path]:
    def strip(im):
        clean = im.copy()
        icc = im.info.get("icc_profile")
        clean.info = {"icc_profile": icc} if icc else {}
        return clean
    return _img_tool(src, "sin-metadatos", strip, keep_meta=False, quality=95)


def imgs_to_pdf(paths: list[Path]) -> list[Path]:
    pages = [_flatten(open_image(p)) for p in paths]
    stem = paths[0].stem if len(paths) == 1 else "Imágenes combinadas"
    out = unique_path(paths[0].parent, stem, "pdf")
    pages[0].save(out, "PDF", save_all=True, append_images=pages[1:], resolution=150)
    return [out]


def imgs_collage(paths: list[Path]) -> list[Path]:
    from PIL import Image
    imgs = [_flatten(open_image(p)) for p in paths]
    cols = math.ceil(math.sqrt(len(imgs)))
    rows = math.ceil(len(imgs) / cols)
    cell, gap = 900, 24
    sheet = Image.new("RGB", (cols * cell + (cols + 1) * gap, rows * cell + (rows + 1) * gap),
                      (255, 255, 255))
    for i, im in enumerate(imgs):
        im = im.copy()
        im.thumbnail((cell, cell), Image.Resampling.LANCZOS)
        r, c = divmod(i, cols)
        x = gap + c * (cell + gap) + (cell - im.width) // 2
        y = gap + r * (cell + gap) + (cell - im.height) // 2
        sheet.paste(im, (x, y))
    out = unique_path(paths[0].parent, "Collage", "jpg")
    sheet.save(out, "JPEG", quality=90, optimize=True)
    return [out]


# ---------------------------------------------------------------- audio y video

AUDIO_ARGS = {
    "mp3": ["-c:a", "libmp3lame", "-q:a", "2"],
    "m4a": ["-c:a", "aac", "-b:a", "192k"],
    "wav": ["-c:a", "pcm_s16le"],
    "flac": ["-c:a", "flac"],
    "ogg": ["-c:a", "libvorbis", "-q:a", "5"],
    "opus": ["-c:a", "libopus", "-b:a", "128k"],
    "aiff": ["-c:a", "pcm_s16be"],
    "wma": ["-c:a", "wmav2", "-b:a", "192k"],
}
_EVEN = "scale=trunc(iw/2)*2:trunc(ih/2)*2"
_H264 = ["-c:v", "libx264", "-preset", "medium", "-crf", "20", "-pix_fmt", "yuv420p"]
VIDEO_ARGS = {
    "mp4": [*_H264, "-vf", _EVEN, "-c:a", "aac", "-b:a", "160k", "-movflags", "+faststart"],
    "mov": [*_H264, "-vf", _EVEN, "-c:a", "aac", "-b:a", "160k", "-movflags", "+faststart"],
    "mkv": [*_H264, "-vf", _EVEN, "-c:a", "aac", "-b:a", "160k"],
    "webm": ["-c:v", "libvpx-vp9", "-crf", "33", "-b:v", "0", "-row-mt", "1", "-deadline", "good",
             "-cpu-used", "4", "-c:a", "libopus", "-b:a", "128k"],
    "avi": ["-c:v", "mpeg4", "-q:v", "4", "-vf", _EVEN, "-c:a", "libmp3lame", "-q:a", "3"],
    "wmv": ["-c:v", "wmv2", "-b:v", "4M", "-c:a", "wmav2", "-b:a", "192k"],
}
_GIF_VF = ("fps=12,scale='min(640,iw)':-2:flags=lanczos,split[a][b];"
           "[a]palettegen=stats_mode=diff[p];[b][p]paletteuse=dither=bayer:bayer_scale=4")


def _audio_args(ext: str) -> list[str]:
    return AUDIO_ARGS.get(ext, AUDIO_ARGS["mp3"])


def audio_convert(src: Path, fmt: str) -> list[Path]:
    out = unique_path(src.parent, src.stem, fmt)
    try:
        ffmpeg(["-i", src, "-vn", "-map_metadata", "0", *_audio_args(fmt), out])
    except GajoError:
        out.unlink(missing_ok=True)
        raise
    return [out]


def video_convert(src: Path, fmt: str) -> list[Path]:
    out = unique_path(src.parent, src.stem, fmt)
    try:
        if fmt == "mp3":
            ffmpeg(["-i", src, "-vn", *_audio_args("mp3"), out])
        elif fmt == "gif":
            ffmpeg(["-i", src, "-vf", _GIF_VF, "-loop", "0", out])
        else:
            info = probe(src)
            can_copy = (fmt in ("mp4", "mov", "mkv") and info["v"] in ("h264", "hevc")
                        and info["a"] in (None, "aac", "mp3"))
            if can_copy:  # mismo contenido, otra "caja": instantaneo y sin perder calidad
                extra = ["-movflags", "+faststart"] if fmt != "mkv" else []
                ffmpeg(["-i", src, "-map", "0:v:0", "-map", "0:a?", "-c", "copy", *extra, out])
            else:
                ffmpeg(["-i", src, "-map", "0:v:0", "-map", "0:a?", *VIDEO_ARGS[fmt], out])
    except GajoError:
        out.unlink(missing_ok=True)
        raise
    return [out]


def _video_out_ext(src: Path) -> str:
    ext = norm_ext(src)
    return ext if ext in ("mp4", "mov", "mkv") else "mp4"


def _ff_tool(src: Path, suffix: str, args: list, ext: str) -> list[Path]:
    out = unique_path(src.parent, f"{src.stem}-{suffix}", ext)
    try:
        ffmpeg(["-i", src, *args, out])
    except GajoError:
        out.unlink(missing_ok=True)
        raise
    return [out]


def video_compress(src: Path) -> list[Path]:
    return _ff_tool(src, "comprimido", [
        "-c:v", "libx264", "-preset", "veryfast", "-crf", "28", "-pix_fmt", "yuv420p",
        "-vf", "scale='min(1920,iw)':-2", "-c:a", "aac", "-b:a", "128k",
        "-movflags", "+faststart"], "mp4")


def video_720(src: Path) -> list[Path]:
    return _ff_tool(src, "720p", [
        *_H264, "-vf", "scale=-2:'min(720,ih)'", "-c:a", "aac", "-b:a", "160k",
        "-movflags", "+faststart"], "mp4")


def video_mute(src: Path) -> list[Path]:
    ext = norm_ext(src)
    return _ff_tool(src, "sin-audio", ["-map", "0:v:0", "-c:v", "copy", "-an"],
                    ext if ext != "gif" else "mp4")


def video_rotate(src: Path) -> list[Path]:
    return _ff_tool(src, "girado", [*_H264, "-vf", "transpose=1", "-c:a", "copy"],
                    _video_out_ext(src))


def video_frame(src: Path) -> list[Path]:
    at = max(0.0, probe(src)["duration"] * 0.1)
    out = unique_path(src.parent, f"{src.stem}-fotograma", "png")
    ffmpeg(["-ss", f"{at:.2f}", "-i", src, "-frames:v", "1", out])
    return [out]


def audio_normalize(src: Path) -> list[Path]:
    ext = norm_ext(src)
    return _ff_tool(src, "normalizado", ["-vn", "-af", "loudnorm=I=-16:TP=-1.5:LRA=11",
                                         *_audio_args(ext)], ext if ext in AUDIO_ARGS else "mp3")


def audio_compress(src: Path) -> list[Path]:
    return _ff_tool(src, "comprimido", ["-vn", "-c:a", "libmp3lame", "-b:a", "128k"], "mp3")


def audio_mono(src: Path) -> list[Path]:
    ext = norm_ext(src)
    return _ff_tool(src, "mono", ["-vn", "-ac", "1", *_audio_args(ext)],
                    ext if ext in AUDIO_ARGS else "mp3")


def audio_trim_silence(src: Path) -> list[Path]:
    edge = "silenceremove=start_periods=1:start_threshold=-50dB:start_silence=0.1"
    ext = norm_ext(src)
    return _ff_tool(src, "sin-silencios", ["-vn", "-af", f"{edge},areverse,{edge},areverse",
                                           *_audio_args(ext)], ext if ext in AUDIO_ARGS else "mp3")


def audio_join(paths: list[Path]) -> list[Path]:
    ext = norm_ext(paths[0])
    ext = ext if ext in AUDIO_ARGS else "mp3"
    out = unique_path(paths[0].parent, f"{paths[0].stem}-unido", ext)
    args: list = []
    for p in paths:
        args += ["-i", p]
    labels = "".join(f"[{i}:a:0]" for i in range(len(paths)))
    args += ["-filter_complex", f"{labels}concat=n={len(paths)}:v=0:a=1[a]", "-map", "[a]",
             *_audio_args(ext), out]
    ffmpeg(args)
    return [out]


# ---------------------------------------------------------------- PDF

def _fitz():
    import fitz
    return fitz


def pdf_to_images(src: Path, fmt: str, dpi: int = 300, stem: str | None = None) -> list[Path]:
    fitz = _fitz()
    stem = stem or src.stem
    with fitz.open(src) as doc:
        if len(doc) == 1:
            out = unique_path(src.parent, stem, fmt)
            _save_pix(doc[0].get_pixmap(dpi=dpi), out, fmt)
            return [out]
        from . import progress
        folder = unique_dir(src.parent, f"{stem} - páginas")
        for i, page in enumerate(doc, 1):
            progress.report((i - 1) / len(doc), f"Página {i} de {len(doc)}")
            _save_pix(page.get_pixmap(dpi=dpi), folder / f"{stem}-{i:03d}.{fmt}", fmt)
        return [folder]


def _save_pix(pix, out: Path, fmt: str) -> None:
    if fmt == "jpg":
        pix.save(out, jpg_quality=92)
    else:
        pix.save(out)


def _pdf2docx(src: Path, out: Path) -> None:
    from pdf2docx import Converter
    for name in ("pdf2docx", "root"):
        logging.getLogger(name).setLevel(logging.WARNING)
    cv = Converter(str(src))
    try:
        cv.convert(str(out))
    finally:
        cv.close()


def pdf_to_docx(src: Path) -> list[Path]:
    out = unique_path(src.parent, src.stem, "docx")
    _pdf2docx(src, out)
    return [out]


def pdf_to_pptx(src: Path) -> list[Path]:
    from pptx import Presentation
    from pptx.util import Emu
    fitz = _fitz()
    prs = Presentation()
    with fitz.open(src) as doc:
        first = doc[0].rect
        prs.slide_width = Emu(int(first.width * 12700))
        prs.slide_height = Emu(int(first.height * 12700))
        blank = prs.slide_layouts[6]
        from . import progress
        for i, page in enumerate(doc):
            progress.report(i / len(doc), f"Página {i + 1} de {len(doc)}")
            png = page.get_pixmap(dpi=200).tobytes("png")
            slide = prs.slides.add_slide(blank)
            slide.shapes.add_picture(io.BytesIO(png), 0, 0, prs.slide_width, prs.slide_height)
    out = unique_path(src.parent, src.stem, "pptx")
    prs.save(out)
    return [out]


_CL_THOUSANDS = re.compile(r"^-?\d{1,3}(\.\d{3})+(,\d+)?$")
_CL_DECIMAL = re.compile(r"^-?\d+(,\d+)?$")
_DOT_DECIMAL = re.compile(r"^-?\d+\.\d+$")


def _cell(value):
    """Convierte textos numericos (formato chileno 1.234,5) en numeros reales."""
    if value is None:
        return ""
    text = str(value).strip()
    if _CL_THOUSANDS.match(text):
        return float(text.replace(".", "").replace(",", "."))
    if _CL_DECIMAL.match(text):
        num = float(text.replace(",", "."))
        return int(num) if num.is_integer() and "," not in text else num
    if _DOT_DECIMAL.match(text):
        return float(text)
    return text


def pdf_to_xlsx(src: Path) -> list[Path]:
    import openpyxl
    fitz = _fitz()
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    with fitz.open(src) as doc:
        for pno, page in enumerate(doc, 1):
            for ti, table in enumerate(page.find_tables().tables, 1):
                ws = wb.create_sheet(f"Pag{pno}-Tabla{ti}"[:31])
                for row in table.extract():
                    ws.append([_cell(c) for c in row])
        if not wb.sheetnames:  # sin tablas: el texto, una linea por fila
            ws = wb.create_sheet("Texto")
            for page in doc:
                for line in page.get_text().splitlines():
                    if line.strip():
                        ws.append([_cell(line)])
    out = unique_path(src.parent, src.stem, "xlsx")
    wb.save(out)
    return [out]


def pdf_to_txt(src: Path) -> list[Path]:
    fitz = _fitz()
    with fitz.open(src) as doc:
        text = "\n\n".join(page.get_text().strip() for page in doc)
    out = unique_path(src.parent, src.stem, "txt")
    out.write_text(text + "\n", encoding="utf-8")
    return [out]


def pdf_to_md(src: Path) -> list[Path]:
    """Texto con titulos detectados por el tamano de la letra."""
    fitz = _fitz()
    from collections import Counter
    with fitz.open(src) as doc:
        pages = [page.get_text("dict") for page in doc]
    sizes: Counter = Counter()
    for pd in pages:
        for block in pd["blocks"]:
            for line in block.get("lines", []):
                for span in line["spans"]:
                    sizes[round(span["size"])] += len(span["text"])
    body = sizes.most_common(1)[0][0] if sizes else 11
    out_lines: list[str] = []
    for pd in pages:
        for block in pd["blocks"]:
            lines = block.get("lines", [])
            text = " ".join("".join(s["text"] for s in ln["spans"]).strip() for ln in lines).strip()
            if not text:
                continue
            size = max((s["size"] for ln in lines for s in ln["spans"]), default=body)
            if size >= body * 1.5 and len(text) < 120:
                out_lines.append(f"# {text}")
            elif size >= body * 1.2 and len(text) < 160:
                out_lines.append(f"## {text}")
            else:
                out_lines.append(text)
            out_lines.append("")
    out = unique_path(src.parent, src.stem, "md")
    out.write_text("\n".join(out_lines), encoding="utf-8")
    return [out]


def pdf_to_epub(src: Path) -> list[Path]:
    tmp = work_dir()
    try:
        docx = tmp / "doc.docx"
        _pdf2docx(src, docx)
        epub = tmp / "doc.epub"
        pandoc([docx, "-o", epub, "--metadata", f"title={src.stem}"])
        return [place(epub, src.parent, src.stem, "epub")]
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def pdf_compress(src: Path) -> list[Path]:
    fitz = _fitz()
    out = unique_path(src.parent, f"{src.stem}-comprimido", "pdf")
    with fitz.open(src) as doc:
        try:
            doc.rewrite_images(dpi_threshold=200, dpi_target=150, quality=70)
        except Exception as exc:  # PDFs raros: igual se limpia y comprime la estructura
            log.warning("rewrite_images fallo en %s: %s", src, exc)
        doc.save(out, garbage=4, deflate=True, deflate_images=True, deflate_fonts=True,
                 clean=True, use_objstms=1)
    return [out]


def pdf_split(src: Path) -> list[Path]:
    fitz = _fitz()
    with fitz.open(src) as doc:
        if len(doc) < 2:
            raise GajoError(f"{src.name} tiene una sola página")
        folder = unique_dir(src.parent, f"{src.stem} - páginas")
        from . import progress
        for i in range(len(doc)):
            progress.report(i / len(doc), f"Página {i + 1} de {len(doc)}")
            with fitz.open() as one:
                one.insert_pdf(doc, from_page=i, to_page=i)
                one.save(folder / f"{src.stem}-{i + 1:03d}.pdf", garbage=3, deflate=True)
    return [folder]


def pdf_rotate(src: Path) -> list[Path]:
    fitz = _fitz()
    out = unique_path(src.parent, f"{src.stem}-girado", "pdf")
    with fitz.open(src) as doc:
        for page in doc:
            page.set_rotation((page.rotation + 90) % 360)
        doc.save(out, garbage=3, deflate=True)
    return [out]


def pdf_extract_images(src: Path) -> list[Path]:
    fitz = _fitz()
    with fitz.open(src) as doc:
        seen: set[int] = set()
        found = []
        for page in doc:
            for img in page.get_images(full=True):
                if img[0] not in seen:
                    seen.add(img[0])
                    found.append(img[0])
        if not found:
            raise GajoError(f"{src.name} no tiene imágenes")
        folder = unique_dir(src.parent, f"{src.stem} - imágenes")
        for n, xref in enumerate(found, 1):
            data = doc.extract_image(xref)
            (folder / f"imagen-{n:03d}.{data['ext']}").write_bytes(data["image"])
    return [folder]


def pdf_merge(paths: list[Path]) -> list[Path]:
    fitz = _fitz()
    out = unique_path(paths[0].parent, f"{paths[0].stem}-unido", "pdf")
    with fitz.open() as merged:
        for p in paths:
            with fitz.open(p) as doc:
                merged.insert_pdf(doc)
        merged.save(out, garbage=3, deflate=True)
    return [out]


# ---------------------------------------------------------------- documentos

_LO_DOC_IN = {"docx", "doc", "odt", "rtf", "txt", "html"}
_PANDOC_IN = {"md": "markdown", "txt": "markdown", "html": "html", "docx": "docx", "odt": "odt",
              "epub": "epub", "rtf": "rtf"}
_PANDOC_OUT = {"docx": "docx", "odt": "odt", "rtf": "rtf", "txt": "plain", "md": "gfm",
               "html": "html", "epub": "epub"}


def _to_pandoc_source(src: Path, tmp: Path) -> tuple[Path, str]:
    ext = norm_ext(src)
    if ext in _PANDOC_IN:
        return src, _PANDOC_IN[ext]
    return soffice_convert(src, "docx", tmp), "docx"   # .doc y otros viejos


def doc_convert(src: Path, fmt: str) -> list[Path]:
    ext = norm_ext(src)
    tmp = work_dir()
    try:
        if fmt == "pdf":
            if ext in _LO_DOC_IN:
                made = soffice_convert(src, "pdf", tmp)
            else:  # md, epub: primero a Word con Pandoc, luego a PDF con LibreOffice
                docx = tmp / f"{src.stem}.docx"
                pandoc([src, "-f", _PANDOC_IN[ext], "-o", docx])
                made = soffice_convert(docx, "pdf", tmp)
        elif fmt in ("docx", "odt") and ext in {"doc", "rtf", "docx", "odt"}:
            made = soffice_convert(src, fmt, tmp)
        else:
            source, reader = _to_pandoc_source(src, tmp)
            made = tmp / f"salida.{fmt}"
            args = [source, "-f", reader, "-t", _PANDOC_OUT[fmt], "-o", made]
            if fmt in ("html", "rtf"):
                args += ["-s", "--embed-resources", "--metadata", f"title={src.stem}"]
            elif fmt == "epub":
                args += ["--metadata", f"title={src.stem}"]
            pandoc(args)
        return [place(made, src.parent, src.stem, fmt)]
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


# ---------------------------------------------------------------- presentaciones y planillas

def slides_convert(src: Path, fmt: str) -> list[Path]:
    tmp = work_dir()
    try:
        if fmt in ("pdf", "pptx"):
            made = soffice_convert(src, fmt, tmp)
            return [place(made, src.parent, src.stem, fmt)]
        pdf = soffice_convert(src, "pdf", tmp)
        return pdf_to_images_into(pdf, fmt, src)
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def pdf_to_images_into(pdf: Path, fmt: str, original: Path) -> list[Path]:
    """Igual que pdf_to_images pero dejando el resultado junto al archivo original."""
    fitz = _fitz()
    with fitz.open(pdf) as doc:
        if len(doc) == 1:
            out = unique_path(original.parent, original.stem, fmt)
            _save_pix(doc[0].get_pixmap(dpi=150), out, fmt)
            return [out]
        folder = unique_dir(original.parent, f"{original.stem} - diapositivas")
        for i, page in enumerate(doc, 1):
            _save_pix(page.get_pixmap(dpi=150), folder / f"{original.stem}-{i:03d}.{fmt}", fmt)
        return [folder]


_SHEET_OUT = {
    "xlsx": "xlsx:Calc MS Excel 2007 XML",
    "ods": "ods",
    "pdf": "pdf",
    "csv": "csv:Text - txt - csv (StarCalc):44,34,76,1",
    "tsv": "csv:Text - txt - csv (StarCalc):9,34,76,1",
}


def sheet_convert(src: Path, fmt: str) -> list[Path]:
    ext = norm_ext(src)
    infilter = {"csv": "CSV:44,34,76,1", "tsv": "CSV:9,34,76,1"}.get(ext)
    tmp = work_dir()
    try:
        made = soffice_convert(src, _SHEET_OUT[fmt], tmp, infilter)
        return [place(made, src.parent, src.stem, fmt)]
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


# ---------------------------------------------------------------- subtitulos

_TIMING = re.compile(
    r"((?:\d+:)?\d{1,2}:\d{2}[.,]\d{3})\s*-->\s*((?:\d+:)?\d{1,2}:\d{2}[.,]\d{3})")


def _read_text(path: Path) -> str:
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "cp1252"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("latin-1")


def _ms(stamp: str) -> int:
    parts = stamp.replace(",", ".").split(":")
    secs = float(parts[-1])
    mins = int(parts[-2])
    hours = int(parts[-3]) if len(parts) == 3 else 0
    return int(round((hours * 3600 + mins * 60 + secs) * 1000))


def _stamp(ms: int, sep: str) -> str:
    h, rest = divmod(ms, 3_600_000)
    m, rest = divmod(rest, 60_000)
    s, ms = divmod(rest, 1000)
    return f"{h:02d}:{m:02d}:{s:02d}{sep}{ms:03d}"


def parse_subs(path: Path) -> list[tuple[int, int, str]]:
    cues = []
    lines = _read_text(path).replace("\r\n", "\n").split("\n")
    i = 0
    while i < len(lines):
        match = _TIMING.search(lines[i])
        if match:
            text = []
            i += 1
            while i < len(lines) and lines[i].strip():
                text.append(lines[i].rstrip())
                i += 1
            cues.append((_ms(match.group(1)), _ms(match.group(2)), "\n".join(text)))
        i += 1
    if not cues:
        raise GajoError(f"No encontré subtítulos en {path.name}")
    return cues


def subs_convert(src: Path, fmt: str) -> list[Path]:
    cues = parse_subs(src)
    if fmt == "srt":
        body = "\n".join(f"{n}\n{_stamp(a, ',')} --> {_stamp(b, ',')}\n{t}\n"
                         for n, (a, b, t) in enumerate(cues, 1))
    elif fmt == "vtt":
        body = "WEBVTT\n\n" + "\n".join(f"{_stamp(a, '.')} --> {_stamp(b, '.')}\n{t}\n"
                                        for a, b, t in cues)
    else:
        body = "\n".join(re.sub(r"<[^>]+>", "", t).replace("\n", " ") for _, _, t in cues) + "\n"
    out = unique_path(src.parent, src.stem, fmt)
    out.write_text(body, encoding="utf-8")
    return [out]


# ---------------------------------------------------------------- archivos comprimidos

def _safe_target(dest: Path, name: str) -> Path:
    target = (dest / name).resolve()
    if dest.resolve() not in target.parents and target != dest.resolve():
        raise GajoError(f"Ruta peligrosa dentro del comprimido: {name}")
    return target


def _extract_to(src: Path, dest: Path) -> None:
    ext = norm_ext(src)
    if ext == "zip":
        with zipfile.ZipFile(src) as zf:
            for info in zf.infolist():
                name = info.filename
                if not info.flag_bits & 0x800:  # nombres viejos de Windows (tildes, enes)
                    try:
                        name = name.encode("cp437").decode("cp850")
                    except UnicodeError:
                        pass
                target = _safe_target(dest, name)
                if info.is_dir():
                    target.mkdir(parents=True, exist_ok=True)
                    continue
                target.parent.mkdir(parents=True, exist_ok=True)
                with zf.open(info) as fin, open(target, "wb") as fout:
                    shutil.copyfileobj(fin, fout)
    elif ext in ("tar", "tgz", "gz", "bz2", "xz") and tarfile.is_tarfile(src):
        with tarfile.open(src) as tf:
            tf.extractall(dest, filter="data")
    elif ext == "gz":
        import gzip
        with gzip.open(src) as fin, open(dest / src.stem, "wb") as fout:
            shutil.copyfileobj(fin, fout)
    else:  # rar, 7z: el tar de Windows (libarchive) los sabe leer
        run([BSDTAR or "tar", "-xf", src, "-C", dest], 1800)


def archive_extract(src: Path) -> list[Path]:
    folder = unique_dir(src.parent, base_stem(src))
    try:
        _extract_to(src, folder)
    except Exception:
        shutil.rmtree(folder, ignore_errors=True)
        raise
    return [folder]


def _zip_items(items: list[Path], out: Path) -> None:
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED, compresslevel=6) as zf:
        for item in items:
            if item.is_dir():
                for f in sorted(item.rglob("*")):
                    arc = f.relative_to(item.parent).as_posix()
                    if f.is_dir():
                        zf.writestr(arc + "/", "")
                    else:
                        zf.write(f, arc)
            else:
                zf.write(item, item.name)


def archive_convert(src: Path, fmt: str) -> list[Path]:
    tmp = work_dir()
    try:
        _extract_to(src, tmp)
        items = sorted(tmp.iterdir())
        stem = base_stem(src)
        made = tmp.parent / f"{tmp.name}.{fmt}"
        if fmt == "zip":
            _zip_items(items, made)
        else:
            with tarfile.open(made, "w:gz" if fmt == "tgz" else "w") as tf:
                for item in items:
                    tf.add(item, arcname=item.name)
        return [place(made, src.parent, stem, "tar.gz" if fmt == "tgz" else fmt)]
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def make_zip(paths: list[Path]) -> list[Path]:
    stem = base_stem(paths[0]) if len(paths) == 1 else "Comprimido"
    out = unique_path(paths[0].parent, stem, "zip")
    _zip_items(paths, out)
    return [out]
